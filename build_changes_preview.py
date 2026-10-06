"""
Genera un Excel de PROPUESTA (antes de ejecutar nada en Shopify) con el detalle de
cada pedido que diff.py detecto con algun cambio pendiente: tags actuales, tags que
quedarian, que se agregaria/quitaria, y el movimiento/split propuesto si aplica.

Requiere que ya se hayan corrido, en este orden, sobre datos recien extraidos:
  1. engine.py   -> scope.pkl
  2. diff.py     -> diff.pkl (imprime el resumen de conteos)

Uso: python3 build_changes_preview.py <fecha DD-MM-AAAA>
Genera: PROPUESTA_<fecha>.xlsx en el directorio actual.

v66+ (21/09/2026): pestana nueva "D10 no activado" con las variantes detectadas por el
paso 2B (verificacion/activacion de D10). Se alimenta de `d10_pending.json` si existe
en el directorio de trabajo (lo escribe el paso 2B); si no existe, la pestana se crea
igual con un aviso, para que el Excel nunca mienta por omision.
"""
import sys, pickle, json, os
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

fecha = sys.argv[1] if len(sys.argv) > 1 else "SIN-FECHA"

with open("scope.pkl", "rb") as f:
    results = pickle.load(f)
with open("diff.pkl", "rb") as f:
    diff = pickle.load(f)

by_name = {r["order"]["name"]: r for r in results}

changes = {}

def entry(name, order_id, created):
    if name not in changes:
        changes[name] = {"id": order_id, "fecha": created, "add": [], "remove": [], "move": None, "split": None}
    return changes[name]

for t in diff["tagsAdd"]:
    r = by_name.get(t["order"])
    created = r["order"]["createdAt"] if r else ""
    e = entry(t["order"], t["id"], created)
    e["add"].extend(t["tags"])

for t in diff["tagsRemove"]:
    r = by_name.get(t["order"])
    created = r["order"]["createdAt"] if r else ""
    e = entry(t["order"], t["id"], created)
    e["remove"].extend(t["tags"])

for m in diff["moves"]:
    r = by_name.get(m["order"])
    created = r["order"]["createdAt"] if r else ""
    e = entry(m["order"], r["order"]["id"] if r else None, created)
    e["move"] = f'{m["from"]} -> {m["to"]} (fulfillment order {m["fo_id"].split("/")[-1]})'

for s in diff["splits"]:
    r = by_name.get(s["order"])
    created = r["order"]["createdAt"] if r else ""
    e = entry(s["order"], r["order"]["id"] if r else None, created)
    partes = []
    for d in s["decisions"]:
        partes.append(f'{d.get("sku") or d.get("name","(sin sku)")} x{d["qty"]} -> {d["destino"]} ({d["motivo"]})')
    e["split"] = " | ".join(partes)

wb = Workbook()
ws = wb.active
ws.title = f"Propuesta {fecha}"
bold = Font(bold=True)
header_fill = PatternFill("solid", fgColor="DDDDDD")

headers = ["Pedido", "Fecha creacion", "Grupo", "Tags ACTUALES", "Tags PROPUESTOS",
           "Tags a agregar", "Tags a quitar", "Movimiento propuesto (FO)", "Split propuesto (FO)"]
for i, h in enumerate(headers, start=1):
    c = ws.cell(row=1, column=i, value=h)
    c.font = bold
    c.fill = header_fill

row = 2
for name in sorted(changes.keys()):
    e = changes[name]
    r = by_name.get(name)
    old_tags = r["order"]["tags"] if r else []
    grupo = r["group"] if r else ""
    add = sorted(set(e["add"]))
    rem = sorted(set(e["remove"]))
    new_tags = sorted((set(old_tags) - set(rem)) | set(add))
    fecha_creacion = (e["fecha"] or "")[:10]
    vals = [name, fecha_creacion, grupo, ", ".join(old_tags), ", ".join(new_tags),
            ", ".join(add), ", ".join(rem), e["move"] or "", e["split"] or ""]
    for i, v in enumerate(vals, start=1):
        ws.cell(row=row, column=i, value=v)
    row += 1

widths = [12, 13, 26, 30, 30, 20, 25, 35, 55]
for col, w in zip("ABCDEFGHI", widths):
    ws.column_dimensions[col].width = w

# --- Pestana ATC: pedidos Pendiente / Parcialmente pagado ---
# Agregada 29/08 a pedido de Northdeco: estos pedidos ya estan fuera del scope de
# almacen/tags (grupo "11_NO_PAGADO" en engine.py, no se tocan ni se cuentan ahi),
# pero Northdeco quiere visibilidad para que el equipo de ATC les haga seguimiento
# del cobro. Se filtra especificamente a estado de pago PENDING/PARTIALLY_PAID (no
# el grupo 11_NO_PAGADO completo, que tambien incluye AUTHORIZED/VOIDED/REFUNDED/
# EXPIRED — esos no son casos de "falta cobrar", asi que no son de ATC).
ATC_STATUSES = {"PENDING", "PARTIALLY_PAID"}
FSTATUS_LABEL = {"PENDING": "Pendiente", "PARTIALLY_PAID": "Parcialmente pagado"}
# sourceName (API) -> nombre de canal como lo muestra el admin de Shopify. "web" es
# el valor que devuelve la API para "Online Store"; se deja el crudo como fallback
# para cualquier canal no mapeado (pos, draft orders, etc.) en vez de ocultarlo.
CHANNEL_LABEL = {"web": "Online Store", "pos": "Point of Sale", "shopify_draft_order": "Draft orders"}

ws_atc = wb.create_sheet("ATC - Pagos pendientes")
atc_headers = ["Pedido", "Etiquetas", "Fecha", "Cliente", "Canal", "Total", "Estado de pago"]
for i, h in enumerate(atc_headers, start=1):
    c = ws_atc.cell(row=1, column=i, value=h)
    c.font = bold
    c.fill = header_fill

atc_rows = [r for r in results if (r["order"].get("financial_status") or "") in ATC_STATUSES]
atc_rows.sort(key=lambda r: r["order"]["createdAt"], reverse=True)

arow = 2
for r in atc_rows:
    o = r["order"]
    fstatus = o.get("financial_status") or ""
    total = o.get("total_amount")
    currency = o.get("total_currency") or ""
    total_str = f"{total} {currency}".strip() if total is not None else ""
    channel_raw = o.get("channel") or ""
    vals = [o["name"], ", ".join(o.get("tags", [])), (o["createdAt"] or "")[:10],
            o.get("customer_name") or "", CHANNEL_LABEL.get(channel_raw, channel_raw),
            total_str, FSTATUS_LABEL.get(fstatus, fstatus)]
    for i, v in enumerate(vals, start=1):
        ws_atc.cell(row=arow, column=i, value=v)
    arow += 1

atc_widths = [12, 30, 11, 26, 14, 14, 20]
for col, w in zip("ABCDEFG", atc_widths):
    ws_atc.column_dimensions[col].width = w

# --- Pestana "D10 no activado" (paso 2B del prompt maestro) ---
# Variantes de producto ACTIVE, con SKU, sin prefijo PZ-, que NO tenian la sucursal D10
# activada en su inventario en el momento de la extraccion de esta corrida. Es la lista
# sobre la que el paso 2B ejecuta `inventoryActivate` (con el tope de seguridad de 25).
# La columna "Activado en esta corrida" deja registro de si se llego a ejecutar o si el
# tope bloqueo la activacion.
ws_d10 = wb.create_sheet("D10 no activado")
d10_headers = ["SKU", "Producto", "Variante", "Ubicaciones que SI tenia",
               "¿Falta tambien Matadepera?", "Activado en esta corrida"]
for i, h in enumerate(d10_headers, start=1):
    c = ws_d10.cell(row=1, column=i, value=h)
    c.font = bold
    c.fill = header_fill

d10_rows = []
if os.path.exists("d10_pending.json"):
    with open("d10_pending.json", encoding="utf-8") as f:
        d10_rows = json.load(f)
else:
    ws_d10.cell(row=2, column=1, value="(no se genero d10_pending.json en esta corrida)")

drow = 2
for v in d10_rows:
    vals = [v.get("sku", ""), v.get("product_title", ""), v.get("variant_title", ""),
            ", ".join(v.get("locations", [])),
            "SI" if v.get("falta_matadepera") else "no",
            v.get("activado", "")]
    for i, val in enumerate(vals, start=1):
        ws_d10.cell(row=drow, column=i, value=val)
    drow += 1

d10_widths = [30, 42, 24, 40, 26, 24]
for col, w in zip("ABCDEF", d10_widths):
    ws_d10.column_dimensions[col].width = w

out_name = f"PROPUESTA_{fecha}.xlsx"
wb.save(out_name)
print("saved:", out_name, "filas:", row - 2, "| ATC pendientes de pago:", arow - 2,
      "| D10 no activado:", len(d10_rows))
